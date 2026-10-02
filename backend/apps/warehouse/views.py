"""
仓库管理视图
"""
import logging
import io
from django.db import transaction
from django.db.models import F
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    Appointment, Reception, AppointmentEvent
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    AppointmentSerializer, AppointmentListSerializer, AppointmentCreateSerializer,
    AppointmentRescheduleSerializer, CheckInSerializer
)

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockIn.objects.select_related(
            'goods', 'operator', 'appointment'
        ).order_by('-stock_in_time')

        goods_id = request.query_params.get('goods')
        appointment_id = request.query_params.get('appointment')
        batch_no = request.query_params.get('batch_no')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)
        if appointment_id:
            queryset = queryset.filter(appointment_id=appointment_id)
        if batch_no:
            queryset = queryset.filter(batch_no__icontains=batch_no)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        records = queryset[start:end]

        serializer = StockInSerializer(records, many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })


class StockOutListView(APIView):
    """出库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


# ==================== 到场预约 ====================

class AppointmentListView(APIView):
    """到场预约列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Appointment.objects.select_related('reception').order_by('expected_start', 'id')

        status = request.query_params.get('status')
        transfer_unit = request.query_params.get('transfer_unit')
        date = request.query_params.get('date')
        if status:
            queryset = queryset.filter(status=status)
        if transfer_unit:
            queryset = queryset.filter(transfer_unit__icontains=transfer_unit)
        if date:
            queryset = queryset.filter(expected_start__date=date)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        appointments = queryset[start:end]

        serializer = AppointmentListSerializer(appointments, many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """创建到场预约"""
        serializer = AppointmentCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        data = serializer.validated_data
        with transaction.atomic():
            appointment = Appointment.objects.create(
                appointment_no=Appointment.generate_appointment_no(),
                created_by=request.user,
                **data
            )
            AppointmentEvent.objects.create(
                appointment=appointment,
                event_type='created',
                actor=request.user,
                detail={
                    'transfer_unit': appointment.transfer_unit,
                    'expected_start': appointment.expected_start.isoformat(),
                    'expected_end': appointment.expected_end.isoformat(),
                    'expected_quantity': appointment.expected_quantity,
                    'handover_person': appointment.handover_person,
                }
            )

        logger.info(
            f"User {request.user.username} created appointment "
            f"{appointment.appointment_no} for {appointment.transfer_unit}"
        )

        return success_response(data=AppointmentSerializer(appointment).data, message='创建成功')


class AppointmentDetailView(APIView):
    """到场预约详情视图（含签到记录、轨迹与入库记录）"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            appointment = Appointment.objects.get(pk=pk)
        except Appointment.DoesNotExist:
            return error_response(message='预约不存在', code=404)

        return success_response(data=AppointmentSerializer(appointment).data)


class AppointmentRescheduleView(APIView):
    """预约改期视图"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            appointment = Appointment.objects.get(pk=pk)
        except Appointment.DoesNotExist:
            return error_response(message='预约不存在', code=404)

        if appointment.status != 'pending':
            return error_response(message='仅待到场的预约可以改期')

        serializer = AppointmentRescheduleSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        data = serializer.validated_data
        old_start = appointment.expected_start
        old_end = appointment.expected_end

        with transaction.atomic():
            appointment.expected_start = data['expected_start']
            appointment.expected_end = data['expected_end']
            appointment.save(update_fields=['expected_start', 'expected_end', 'updated_at'])
            AppointmentEvent.objects.create(
                appointment=appointment,
                event_type='rescheduled',
                actor=request.user,
                detail={
                    'old_start': old_start.isoformat(),
                    'old_end': old_end.isoformat(),
                    'new_start': appointment.expected_start.isoformat(),
                    'new_end': appointment.expected_end.isoformat(),
                    'reason': data.get('reason', ''),
                }
            )

        logger.info(
            f"User {request.user.username} rescheduled appointment {appointment.appointment_no}"
        )

        return success_response(data=AppointmentSerializer(appointment).data, message='改期成功')


class AppointmentCancelView(APIView):
    """预约取消视图"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            appointment = Appointment.objects.get(pk=pk)
        except Appointment.DoesNotExist:
            return error_response(message='预约不存在', code=404)

        if appointment.status != 'pending':
            return error_response(message='仅待到场的预约可以取消')

        with transaction.atomic():
            appointment.status = 'cancelled'
            appointment.save(update_fields=['status', 'updated_at'])
            AppointmentEvent.objects.create(
                appointment=appointment,
                event_type='cancelled',
                actor=request.user,
                detail={'reason': request.data.get('reason', '')}
            )

        logger.info(
            f"User {request.user.username} cancelled appointment {appointment.appointment_no}"
        )

        return success_response(data=AppointmentSerializer(appointment).data, message='取消成功')


class AppointmentCheckInView(APIView):
    """到场签到视图：值班员核对身份与数量并决定接收、部分接收或拒收"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        with transaction.atomic():
            try:
                appointment = Appointment.objects.select_for_update().get(pk=pk)
            except Appointment.DoesNotExist:
                return error_response(message='预约不存在', code=404)

            if appointment.status != 'pending':
                reception = getattr(appointment, 'reception', None)
                if reception is not None:
                    # 重复签到：拒绝并保留轨迹
                    AppointmentEvent.objects.create(
                        appointment=appointment,
                        event_type='duplicate_check_in',
                        actor=request.user,
                        detail={
                            'attempted_at': timezone.now().isoformat(),
                            'original_check_in_time': reception.check_in_time.isoformat(),
                            'appointment_status': appointment.status,
                        }
                    )
                    return error_response(message='该预约已完成签到，请勿重复签到')
                return error_response(message='预约已取消，无法签到')

            serializer = CheckInSerializer(data=request.data)
            if not serializer.is_valid():
                errors = serializer.errors
                first_error = list(errors.values())[0]
                if isinstance(first_error, list):
                    first_error = first_error[0]
                return error_response(message=str(first_error))

            data = serializer.validated_data
            now = timezone.now()
            arrival_status = appointment.arrival_status_at(now)
            decision = data['decision']

            reception = Reception.objects.create(
                appointment=appointment,
                duty_officer=request.user,
                check_in_time=now,
                arrival_status=arrival_status,
                actual_person=data['actual_person'],
                actual_person_id_card=data.get('actual_person_id_card', ''),
                identity_verified=data['identity_verified'],
                actual_quantity=data['actual_quantity'],
                decision=decision,
                received_quantity=data['received_quantity'],
                discrepancy_note=data.get('discrepancy_note', ''),
            )

            status_map = {'receive': 'received', 'partial': 'partial', 'reject': 'rejected'}
            appointment.status = status_map[decision]
            appointment.save(update_fields=['status', 'updated_at'])

            AppointmentEvent.objects.create(
                appointment=appointment,
                event_type='check_in',
                actor=request.user,
                detail={
                    'check_in_time': now.isoformat(),
                    'arrival_status': arrival_status,
                    'expected_start': appointment.expected_start.isoformat(),
                    'expected_end': appointment.expected_end.isoformat(),
                    'actual_person': reception.actual_person,
                    'identity_verified': reception.identity_verified,
                    'expected_quantity': appointment.expected_quantity,
                    'actual_quantity': reception.actual_quantity,
                }
            )

            # 接收或部分接收时生成正式入库记录并更新库存
            stock_in_ids = []
            if decision in ('receive', 'partial'):
                for item in data.get('stock_ins', []):
                    stock_in = StockIn.objects.create(
                        goods_id=item['goods'],
                        operator=request.user,
                        quantity=item['quantity'],
                        batch_no=item.get('batch_no', ''),
                        supplier=appointment.transfer_unit,
                        appointment=appointment,
                        remark=item.get('remark', ''),
                    )
                    Goods.objects.filter(pk=item['goods']).update(
                        quantity=F('quantity') + item['quantity']
                    )
                    stock_in_ids.append(stock_in.id)

            event_map = {
                'receive': 'received',
                'partial': 'partial_received',
                'reject': 'rejected',
            }
            AppointmentEvent.objects.create(
                appointment=appointment,
                event_type=event_map[decision],
                actor=request.user,
                detail={
                    'decision': decision,
                    'expected_quantity': appointment.expected_quantity,
                    'actual_quantity': reception.actual_quantity,
                    'received_quantity': reception.received_quantity,
                    'discrepancy_note': reception.discrepancy_note,
                    'stock_in_ids': stock_in_ids,
                }
            )

        logger.info(
            f"User {request.user.username} checked in appointment "
            f"{appointment.appointment_no}: {arrival_status}, decision={decision}"
        )

        return success_response(data=AppointmentSerializer(appointment).data, message='签到完成')
